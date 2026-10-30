"""One-time extraction of templates from the supplied, completed reference proposal.

The originals are never modified. Both templates keep the original design and formatting;
only client-specific values are swapped for {{tokens}} that app/generate.py fills in.
Run outside the production image (needs python-docx and python-pptx):

    python scripts/prepare_templates.py --document "TeamLogicIT - Managed IT Services-ProjectCenter.docx" \
        --presentation ProjectCenter_IT_Proposal.pptx --client-name ProjectCenter --proposal-month "October 2026"

It is tailored to the reference structure (cover-page table positions and slide shape ids).
Inspect the output for leftover client data before committing.
"""
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root))
from app.textfill import paragraphs, replace, replace_all, replace_regex, text  # noqa: E402

out = root / 'templates'
DATE = r'\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{4}'
# Shape ids on the reference slides that hold client-specific values: {slide index: {shape id: (old, token)}}.
SLIDE_VALUES = {
    10: {8: None},  # Current-state bullets, handled separately.
    11: {3: ('4 Devices', '{{pricing_basis}}'), 7: ('$410.00', '{{monthly}}'),
         25: ('$320.00', '{{core_monthly}}'), 27: ('4  Devices × $80.00', '{{core_detail}}'), 28: ('$320.00', '{{core_total}}'),
         70: ('$500', '{{onboarding}}'),
         80: ('$50.00', '{{security_monthly}}'), 35: ('2 users × $25.00', '{{security_detail}}'), 32: ('$50.00', '{{security_total}}'),
         31: ('Google Cloud', '{{backup_label}}'), 16: ('$40.00', '{{backup_monthly}}'),
         6: ('2 users × $20.00', '{{backup_detail}}'), 19: ('$40.00', '{{backup_total}}')},
    12: {7: ('$135', '{{hourly_rate}}'), 12: ('$202', '{{after_hours_rate}}')},
}
LEAK_PATTERNS = [r'\$\d', r'\d+\s+users?\s+×']


def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from all_shapes(shape.shapes)


def set_shape_value(shape, old, token):
    hits = sum(replace(p, old, token) for p in paragraphs(shape._element))
    if hits != 1:
        raise SystemExit(f'Expected one "{old}" in shape {shape.shape_id} ({shape.name}); found {hits}.')


def prepare_presentation(path, client, month):
    prs = Presentation(path)
    for index, slide in enumerate(prs.slides):
        if slide.has_notes_slide:
            slide.notes_slide.notes_text_frame.text = ''
        by_id = {s.shape_id: s for s in all_shapes(slide.shapes)}
        for shape_id, change in SLIDE_VALUES.get(index, {}).items():
            shape = by_id[shape_id]
            if change:
                set_shape_value(shape, *change)
            else:  # Bulleted list: keep the first bullet as the token paragraph, plus the trailing empty one.
                paras = shape.text_frame.paragraphs
                for p in paras[1:]:
                    if p.text.strip():
                        p._p.getparent().remove(p._p)
                replace(paras[0]._p, paras[0].text, '{{current_state}}')
        for shape in list(all_shapes(slide.shapes)):
            if shape.has_text_frame:
                for p in paragraphs(shape._element):
                    replace(p, client, '{{client}}')
                    replace(p, month, '{{proposal_month}}')
                # Slide numbers ("11 / 14") are removed from every slide.
                if re.fullmatch(r'\s*\d+\s*/\s*\d+\s*', shape.text_frame.text):
                    shape._element.getparent().remove(shape._element)
    prs.core_properties.title = 'TeamLogic IT proposal template'
    prs.core_properties.author = 'TeamLogic IT Scottsdale'
    prs.core_properties.last_modified_by = 'TeamLogic IT Scottsdale'
    target = out / 'proposal.pptx'
    prs.save(target)
    for index in SLIDE_VALUES:
        slide_text = ' '.join(s.text_frame.text for s in all_shapes(Presentation(target).slides[index].shapes) if s.has_text_frame)
        for pattern in LEAK_PATTERNS:
            if re.search(pattern, slide_text):
                raise SystemExit(f'Slide {index + 1} still contains a priced value ({pattern}).')


def cell_paragraphs(table, row, col):
    return table.cell(row, col).paragraphs


def prepare_document(path, client):
    doc = Document(path)
    body = doc.element.body
    cover, schedule_c, equipment = doc.tables[0], doc.tables[4], doc.tables[5]
    replace_all(body, client, '{{client}}')
    # Cover page (content control with a text box and its fallback copy).
    for p in paragraphs(body[0]):
        replace_regex(p, DATE, '{{proposal_date}}')
    # SOW cover table.
    address = cell_paragraphs(cover, 1, 1)
    replace(address[1]._p, address[1].text, '{{address}}')
    for p in address[2:4]:
        p._p.getparent().remove(p._p)
    replace(address[4]._p, address[4].text, '{{phone}}')
    replace(cell_paragraphs(cover, 2, 1)[0]._p, cell_paragraphs(cover, 2, 1)[0].text, '{{contact}}')
    replace(cell_paragraphs(cover, 5, 1)[0]._p, cell_paragraphs(cover, 5, 1)[0].text, '{{prepared_by}}')
    replace_regex(cell_paragraphs(cover, 10, 1)[1]._p, DATE, '{{msa_date}}')
    replace_regex(cell_paragraphs(cover, 11, 1)[0]._p, DATE, '{{effective_date}}')
    term = cell_paragraphs(cover, 12, 1)
    replace_regex(term[0]._p, r'(?<=– ).+$', '{{initial_term}}')
    replace_regex(term[1]._p, r'(?<=– ).+$', '{{renewal_term}}')
    replace_regex(term[2]._p, r'\d+(?= days)', '{{notice_days}}')
    fees = cell_paragraphs(cover, 13, 1)
    replace_regex(fees[0]._p, r'\$[\d,.]+', '{{setup_fee}}')
    replace_regex(fees[1]._p, r'\$[\d,.]+', '{{monthly_fee}}')
    replace_regex(fees[1]._p, r'(?<=\().+?(?=\))', '{{fee_basis}}')
    replace_regex(fees[2]._p, r'(?<=\().+?(?=\))', '{{fee_components}}')
    # Schedule C rates and covered equipment.
    rates = cell_paragraphs(schedule_c, 5, 1)
    replace_regex(rates[0]._p, r'\$[\d,.]+', '{{hourly_rate}}')
    replace_regex(rates[1]._p, r'\$[\d,.]+', '{{after_hours_rate}}')
    for row, token in enumerate(['{{workstations}}', '{{networks}}', '{{servers}}', '{{mobile_devices}}']):
        replace_regex(cell_paragraphs(equipment, row, 2)[0]._p, r'\d+', token, count=1)
    doc.core_properties.title = 'TeamLogic IT managed services agreement template'
    doc.core_properties.author = 'TeamLogic IT Scottsdale'
    doc.core_properties.last_modified_by = 'TeamLogic IT Scottsdale'
    doc.save(out / 'agreement.docx')
    services = []
    for table in doc.tables[2:4]:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells[-1] in ('YES', 'NO', 'NA'):
                services.append(dict(description=cells[0], frequency=cells[1],
                                     status={'YES': 'Included', 'NO': 'Excluded', 'NA': 'Not applicable'}[cells[-1]]))
    (out / 'services.json').write_text(json.dumps(services, indent=2), encoding='utf-8')


def check_no_client_data(markers):
    for name in ['proposal.pptx', 'agreement.docx']:
        with zipfile.ZipFile(out / name) as z:
            for part in z.namelist():
                data = z.read(part)
                for marker in markers:
                    if marker and marker.encode() in data:
                        raise SystemExit(f'{name}:{part} still contains "{marker}".')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--document', required=True)
    parser.add_argument('--presentation', required=True)
    parser.add_argument('--client-name', required=True)
    parser.add_argument('--proposal-month', required=True)
    parser.add_argument('--private-marker', action='append', default=[],
                        help='Other client values (contact, street, phone) that must not survive. Repeatable.')
    args = parser.parse_args()
    out.mkdir(exist_ok=True)
    prepare_presentation(args.presentation, args.client_name, args.proposal_month)
    prepare_document(args.document, args.client_name)
    check_no_client_data([args.client_name] + args.private_marker)
    print('Templates written to', out)
