import hashlib
import io
import json
import os
import tempfile
import zipfile
from decimal import Decimal
os.environ['DATA_DIR']=tempfile.mkdtemp(prefix='tlit-test-')
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from app.main import app,connect
from app.models import Proposal,quote
from app.generate import TEMPLATES,docx_bytes,pptx_bytes

def example():return Proposal(**json.loads((TEMPLATES/'example.json').read_text()))

def docx_text(raw):
    """All text in the document, one line per paragraph (tables and text boxes included)."""
    body=Document(io.BytesIO(raw)).element.body
    return '\n'.join(''.join(t.text or '' for t in p.iter() if t.tag.endswith('}t')) for p in body.iter() if p.tag.endswith('}p'))

def deck_text(raw):
    def walk(shapes):
        for s in shapes:
            if s.shape_type==6:yield from walk(s.shapes)
            elif s.has_text_frame:yield s.text_frame.text
    return [' '.join(walk(slide.shapes)) for slide in Presentation(io.BytesIO(raw)).slides]

def test_example_totals():
    q=quote(example())
    assert q['monthly']==Decimal('410.00')
    assert q['one_time']==Decimal('500.00')
    assert q['first_month']==Decimal('910.00')
    assert q['term_total']==Decimal('15260.00')

def test_billing_modes_and_tax_allocation():
    p=example();p.user_rate=Decimal('100');p.billing_mode='users';p.security_quantity=p.backup_quantity=0;p.user_taxable=True
    p.tax_percent=Decimal('8');p.discount=Decimal('20')
    assert quote(p)['monthly']==Decimal('194.40')
    p.billing_mode='combined';p.workstation_taxable=False;p.server_taxable=False
    q=quote(p)
    assert q['subtotal']==Decimal('520.00')
    assert q['tax']==Decimal('15.38')

def test_invalid_discount_and_fractional_counts():
    import pytest
    p=example();p.discount=Decimal('999')
    with pytest.raises(ValueError):quote(p)
    with pytest.raises(ValueError):Proposal(**{**example().model_dump(),'users':1.2})
    with pytest.raises(ValueError):Proposal(**{**example().model_dump(),'user_rate':'NaN'})

def test_auth_save_reload_and_download():
    salt=b'0123456789abcdef';password='test-only-password'
    with connect() as db:
        db.execute('INSERT OR REPLACE INTO users VALUES(?,?,?)',('tester',salt.hex(),hashlib.scrypt(password.encode(),salt=salt,n=16384,r=8,p=1).hex()))
    with TestClient(app) as c:
        assert c.get('/api/proposals').status_code==401
        assert c.post('/api/login',json={'username':'tester','password':password}).status_code==403
        c.headers['X-TLIT-Request']='1'
        assert c.post('/api/login',json={'username':'tester','password':password}).status_code==200
        data=example().model_dump(mode='json')
        result=c.post('/api/proposals',json=data);assert result.status_code==200
        assert c.get('/api/proposals/'+result.json()['id']).json()==data
        data['client']='Different Client & Co'
        assert c.put('/api/proposals/'+result.json()['id'],json=data,headers={'If-Match':'1'}).status_code==200
        response=c.post('/api/generate',json=data);assert response.status_code==200
        z=zipfile.ZipFile(io.BytesIO(response.content))
        assert sorted(z.namelist())==sorted(['Different Client & Co_IT_Proposal.pptx','TeamLogicIT - Managed IT Services-Different Client & Co.docx','proposal-data.json','generation-manifest.json'])
        text=docx_text(z.read('TeamLogicIT - Managed IT Services-Different Client & Co.docx'))
        assert 'Different Client & Co' in text and '“Monthly Fee” – $410.00 (Workstation/Server - 04)' in text
        assert '“Initial Term” – three (3) years' in text and '(Managed Services, Cybersecurity, Backup)' in text
        assert 'ProjectCenter' not in text and '{{' not in text
        assert 'Recurring Payment Authorization Form' in text
        assert 'ninety (90) days' in ' '.join(text.split())
        slides=deck_text(z.read('Different Client & Co_IT_Proposal.pptx'))
        all_text=' '.join(slides)
        assert len(slides)==15 and 'Inventory Recap' in slides[14]
        assert 'ProjectCenter' not in all_text
        assert '{{' not in all_text
        assert '$410.00' in slides[11] and '4  Devices × $80.00' in slides[11] and '2 users × $25.00' in slides[11]
        assert '$135' in slides[12] and '$202' in slides[12]
        import re
        assert not re.search(r'\b\d+ / 1[45]\b',all_text)
        assert c.post('/api/logout').status_code==200
        assert c.get('/api/proposals').status_code==401

def slide_signature(slide):
    """Geometry and text of every shape, including those inside groups."""
    def walk(shapes):
        for s in shapes:
            yield (s.shape_id,s.left,s.top,s.width,s.height,s.text_frame.text if s.has_text_frame else '')
            if s.shape_type==6:yield from walk(s.shapes)
    return sorted(walk(slide.shapes),key=str)

def test_example_reproduces_reference_deck():
    """Slides 1-14 generated from example.json must match the sanitized copy of the completed reference deck."""
    p=example();deck=Presentation(io.BytesIO(pptx_bytes(p,quote(p))))
    reference=Presentation(TEMPLATES/'reference'/'Example_IT_Proposal.pptx')
    assert len(reference.slides)==14 and len(deck.slides)==15
    for index,(made,expected) in enumerate(zip(deck.slides,reference.slides),start=1):
        assert slide_signature(made)==slide_signature(expected),f'slide {index} differs from the reference'

def test_inventory_recap_slides():
    from app.models import Asset
    p=example();p.networks=2;p.mobile_devices=4
    slides=deck_text(pptx_bytes(p,quote(p)))
    assert 'Inventory Recap' in slides[14] and 'EQUIPMENT IN SCOPE' in slides[14]
    for label in ['USERS','WORKSTATIONS','SERVERS','NETWORKS','MOBILE']:assert label in slides[14]
    table=[s for s in Presentation(io.BytesIO(pptx_bytes(p,quote(p)))).slides[14].shapes if s.has_table][0].table
    assert [c.text for c in table.rows[1].cells]==['Workstation','Desktops / laptops','—','3']
    p.assets=[Asset(kind='Workstation',description=f'Laptop {i}',quantity=1) for i in range(15)]
    deck=Presentation(io.BytesIO(pptx_bytes(p,quote(p))))
    cells=lambda slide:[c.text for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells]
    assert len(deck.slides)==16 and 'Inventory Recap (continued)' in deck_text(pptx_bytes(p,quote(p)))[15]
    assert 'Laptop 5' in cells(deck.slides[14]) and 'Laptop 14' in cells(deck.slides[15])

def test_older_drafts_move_charges_into_cards():
    data=example().model_dump(mode='json')
    for key in [k for k in data if k.startswith(('security_','backup_'))]:del data[key]
    data['extras']=[dict(description='Cybersecurity protection',quantity=2,rate='25'),dict(description='Cloud-to-cloud backup',quantity=3,rate='20',category='backup'),
                    dict(description='Printer lease',quantity=1,rate='40')]
    p=Proposal(**data)
    assert (p.security_quantity,p.security_rate,p.backup_quantity,p.backup_rate)==(2,Decimal('25'),3,Decimal('20'))
    assert [x.description for x in p.extras]==['Printer lease']
    assert any('Printer lease' in w for w in quote(p)['warnings'])

def test_long_content_keeps_reference_layout():
    p=example();p.client='A long client company name that should wrap cleanly in the proposal';p.needs='Long discovery note. '*90
    from app.models import Line
    p.extras=[Line(description='Additional support service '+str(i),quantity=2,rate='10') for i in range(30)]
    q=quote(p)
    assert any('not itemised on the pricing slide' in w for w in q['warnings'])
    assert len(deck_text(pptx_bytes(p,q)))==15
    assert p.client in docx_text(docx_bytes(p,q))

def test_pricing_cards_and_schedule_b():
    from app.models import Service
    p=example();p.billing_mode='combined';p.users=12;p.user_rate=Decimal('45');p.current_state='Aging firewall\nNo tested backups'
    p.security_quantity=0;p.backup_quantity=3;p.backup_unit='devices';p.backup_rate=Decimal('30')
    p.services=[Service(description='Manage Servers',frequency='Weekly',status='Excluded'),Service(description='Quarterly vCIO review',frequency='Quarterly')]
    q=quote(p);slides=deck_text(pptx_bytes(p,q))
    assert 'Aging firewall' in slides[10] and 'No tested backups' in slides[10]
    assert '12 Users / 4 Devices' in slides[11] and '3 devices × $30.00' in slides[11]
    assert 'PROTECTION' not in slides[11]  # Card removed when there is no cybersecurity charge.
    doc=Document(io.BytesIO(docx_bytes(p,q)))
    rows=[[c.text.strip() for c in r.cells] for t in doc.tables[2:4] for r in t.rows]
    services=[(r[0],r[-1]) for r in rows if r[-1] in ('YES','NO','NA')]
    assert services==[('Manage Servers','NO'),('Quarterly vCIO review','YES')]
    assert ['Additional Services'] in [list(dict.fromkeys(r)) for r in rows]
    assert ['General'] not in [list(dict.fromkeys(r)) for r in rows]
